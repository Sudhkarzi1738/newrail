import csv
from collections.abc import Iterator
from pathlib import Path

from openpyxl import load_workbook

from app.services.ingestion.profiles import Profile


def _norm(h: object) -> str:
    return " ".join(str(h or "").strip().lower().replace("\n", " ").split())


def _raw_rows(path: Path) -> Iterator[list]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]  # TODO: some ECI workbooks keep data on a named sheet - make configurable
        for row in ws.iter_rows(values_only=True):
            yield list(row)
    elif path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            yield from csv.reader(f)
    else:
        raise ValueError(f"Unsupported file type: {path.suffix}")  # PDFs: see README (extract tables to CSV first)


def read_mapped(path: Path, profile: Profile) -> Iterator[dict]:
    """Finds the header row (first row containing all required fields) and yields canonical-key dicts."""
    header_idx: dict[str, int] | None = None
    for row in _raw_rows(path):
        if header_idx is None:
            cells = [_norm(c) for c in row]
            found = {}
            for field, names in profile.aliases.items():
                for n in names:
                    if n in cells:
                        found[field] = cells.index(n)
                        break
            if all(r in found for r in profile.required):
                header_idx = found
            continue
        if not any(c not in (None, "") for c in row):
            continue
        yield {f: (row[i] if i < len(row) else None) for f, i in header_idx.items()}
    if header_idx is None:
        raise ValueError(f"No header row matched profile '{profile.name}'. Check aliases in profiles.py.")
