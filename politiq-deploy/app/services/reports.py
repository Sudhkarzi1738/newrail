from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table

from app.core.config import AI_DISCLAIMER

_HEAD = ["Year", "House", "State", "Constituency", "Party", "Votes", "Vote %", "Position", "Won", "Margin", "Deposit forfeited"]


def _row(e: dict) -> list:
    f = lambda v: float(v) if v is not None else None  # noqa: E731
    return [e["year"], e["house"].value, e["state"], e["constituency"], e["party"], e["votes"],
            f(e["vote_share_pct"]), e["position"], "Yes" if e["is_winner"] else "No", e["margin"],
            {True: "Yes", False: "No", None: ""}[e["deposit_forfeited"]]]


def candidate_xlsx(career: dict, ai_sections: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Career"
    ws.append([f"Candidate: {career['name']}  |  Contested: {career['contested']}  Won: {career['won']}  Lost: {career['lost']}"])
    ws.append(_HEAD)
    for c in ws[2]:
        c.font = Font(bold=True)
    for e in career["entries"]:
        ws.append(_row(e))
    ws.append([])
    ws.append(["Source: Election Commission of India / MyNeta-ADR as loaded into this platform."])
    if ai_sections:
        a = wb.create_sheet("AI analysis")
        a.append([AI_DISCLAIMER])
        for s in ai_sections:
            a.append([s["kind"], str(s.get("content") or s.get("content_text"))[:30000]])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def candidate_pdf(career: dict, ai_sections: list[dict]) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title=f"Report - {career['name']}")
    st = getSampleStyleSheet()
    story = [Paragraph(f"Candidate report: {career['name']}", st["Title"]),
             Paragraph(f"Contested {career['contested']} | Won {career['won']} | Lost {career['lost']}", st["Normal"]),
             Spacer(1, 12)]
    data = [_HEAD] + [[("" if v is None else str(v)) for v in _row(e)] for e in career["entries"]]
    story.append(Table(data, repeatRows=1, style=[("FONTSIZE", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), 0.25, "grey")]))
    for s in ai_sections:
        story += [Spacer(1, 14), Paragraph(f"{s['kind']} ({AI_DISCLAIMER})", st["Heading3"]),
                  Paragraph(str(s.get("content") or s.get("content_text"))[:6000].replace("<", "&lt;"), st["BodyText"])]
    doc.build(story)
    return buf.getvalue()
